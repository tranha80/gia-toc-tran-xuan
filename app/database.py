import sqlite3
import json
import os
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
DB_PATH = os.path.join(BASE_DIR, "gia_toc_tran_xuan.db")
JSON_PATH = os.path.join(BASE_DIR, "gia_toc_tran_xuan.json")
EXCEL_PATH = os.path.join(BASE_DIR, "Gia_Pha_Tran_Xuan_So_Hoa_Chuan.xlsx")

DEFAULT_RELATIONS = {
    "TX-003": ("TX-001", "", "TX-004", "Vợ cả"),
    "TX-004": ("", "", "TX-003", "Vợ cả"),
    "TX-005": ("", "", "TX-003", "Vợ hai"),
    "TX-006": ("TX-001", "", "", ""),
    "TX-007": ("TX-003", "", "TX-008", "Đang kết hôn"),
    "TX-008": ("", "", "TX-007", "Đang kết hôn"),
    "TX-009": ("TX-007", "TX-008", "", ""),
    "TX-010": ("", "", "TX-011", "Đã ly hôn"),
    "TX-011": ("TX-007", "TX-008", "TX-012", "Vợ hai"),
    "TX-012": ("", "", "TX-011", "Vợ hai"),
    "TX-013": ("TX-007", "TX-008", "", ""),
    "TX-014": ("TX-007", "TX-008", "", ""),
    "TX-015": ("TX-011", "TX-010", "TX-016", "Đang kết hôn"),
    "TX-016": ("", "", "TX-015", "Đang kết hôn"),
    "TX-017": ("TX-011", "TX-012", "TX-018", "Đang kết hôn"),
    "TX-018": ("", "", "TX-017", "Đang kết hôn"),
    "TX-019": ("TX-011", "TX-012", "", ""),
    "TX-020": ("TX-011", "TX-012", "TX-021", "Đang kết hôn"),
    "TX-021": ("", "", "TX-020", "Đang kết hôn"),
    "TX-022": ("TX-011", "TX-012", "", ""),
    "TX-023": ("TX-011", "TX-012", "", ""),
    "TX-024": ("TX-011", "TX-012", "", ""),
    "TX-025": ("TX-011", "TX-012", "", ""),
    "TX-026": ("TX-011", "TX-012", "TX-027", "Đang kết hôn"),
    "TX-027": ("", "", "TX-026", "Đang kết hôn"),
    "TX-028": ("TX-015", "TX-016", "", ""),
    "TX-029": ("TX-015", "TX-016", "TX-030", "Đang kết hôn"),
    "TX-030": ("", "", "TX-029", "Đang kết hôn"),
    "TX-031": ("TX-015", "TX-016", "TX-032", "Đang kết hôn"),
    "TX-032": ("", "", "TX-031", "Đang kết hôn"),
    "TX-033": ("TX-017", "TX-018", "TX-034", "Đã ly hôn"),
    "TX-034": ("", "", "TX-033", "Đã ly hôn"),
    "TX-035": ("", "", "TX-033", "Đã ly hôn"),
    "TX-036": ("TX-017", "TX-018", "TX-038", "Đang kết hôn"),
    "TX-037": ("", "", "TX-036", "Đã ly hôn"),
    "TX-038": ("", "", "TX-036", "Đang kết hôn"),
    "TX-039": ("TX-017", "TX-018", "", ""),
    "TX-040": ("TX-020", "TX-021", "", ""),
    "TX-041": ("TX-020", "TX-021", "TX-042", "Đang kết hôn"),
    "TX-042": ("", "", "TX-041", "Đang kết hôn"),
    "TX-043": ("TX-026", "TX-027", "", ""),
    "TX-044": ("TX-026", "TX-027", "", ""),
    "TX-045": ("TX-029", "TX-030", "", ""),
    "TX-046": ("TX-029", "TX-030", "", ""),
    "TX-047": ("TX-031", "TX-032", "", ""),
    "TX-048": ("TX-031", "TX-032", "", ""),
    "TX-049": ("TX-031", "TX-032", "", ""),
    "TX-050": ("TX-033", "TX-034", "", ""),
    "TX-051": ("TX-033", "TX-034", "", ""),
    "TX-052": ("TX-036", "TX-037", "", ""),
    "TX-053": ("TX-036", "TX-038", "", ""),
    "TX-054": ("TX-041", "TX-042", "", ""),
    "TX-055": ("TX-041", "TX-042", "", "")
}

def ensure_db_schema(conn):
    try:
        c = conn.cursor()
        c.execute("PRAGMA table_info(thanh_vien)")
        existing_cols = [r[1] for r in c.fetchall()]
        new_cols = [
            ("cha_id", "TEXT"),
            ("me_id", "TEXT"),
            ("vo_chong_id", "TEXT"),
            ("tinh_trang_hn", "TEXT")
        ]
        for col_name, col_type in new_cols:
            if col_name not in existing_cols:
                c.execute(f"ALTER TABLE thanh_vien ADD COLUMN {col_name} {col_type}")
        conn.commit()

        # Tự động gán mối quan hệ ban đầu nếu chưa có
        c.execute("SELECT COUNT(*) FROM thanh_vien WHERE (cha_id IS NOT NULL AND cha_id != '') OR (vo_chong_id IS NOT NULL AND vo_chong_id != '')")
        count_rel = c.fetchone()[0]
        if count_rel == 0:
            for mid, (cid, mid_m, vcid, hn) in DEFAULT_RELATIONS.items():
                c.execute("UPDATE thanh_vien SET cha_id = ?, me_id = ?, vo_chong_id = ?, tinh_trang_hn = ? WHERE id = ?", (cid, mid_m, vcid, hn, mid))
            conn.commit()
    except Exception as e:
        print(f"[SCHEMA INIT]: {e}")

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    ensure_db_schema(conn)
    return conn

def get_all_members(query=None, gen=None, gender=None):
    conn = get_db()
    c = conn.cursor()
    sql = "SELECT * FROM thanh_vien WHERE 1=1"
    params = []
    if gen and gen != "all":
        sql += " AND the_he = ?"
        params.append(int(gen))
    if gender and gender != "all":
        sql += " AND gioi_tinh = ?"
        params.append(gender)
    if query:
        sql += " AND (ho_ten LIKE ? OR ghi_chu LIKE ? OR sinh LIKE ? OR vai_ve LIKE ?)"
        q_wild = f"%{query}%"
        params.extend([q_wild, q_wild, q_wild, q_wild])
    sql += " ORDER BY the_he ASC, id ASC"
    c.execute(sql, params)
    rows = [dict(r) for r in c.fetchall()]
    conn.close()
    return rows

def get_member_by_id(member_id):
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT * FROM thanh_vien WHERE id = ?", (member_id,))
    row = c.fetchone()
    conn.close()
    return dict(row) if row else None

def get_next_member_id():
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT id FROM thanh_vien WHERE id LIKE 'TX-%' ORDER BY id DESC")
    rows = c.fetchall()
    conn.close()
    max_num = 0
    for r in rows:
        try:
            num = int(r["id"].replace("TX-", ""))
            if num > max_num:
                max_num = num
        except ValueError:
            pass
    return f"TX-{(max_num + 1):03d}"

def create_member(data):
    conn = get_db()
    c = conn.cursor()
    mid = data.get("id") or get_next_member_id()
    c.execute("""
        INSERT INTO thanh_vien (id, ho_ten, gioi_tinh, the_he, vai_ve, sinh, mat_duong, mat_am, huong_tho, que_quan, phan_mo, ghi_chu, cha_id, me_id, vo_chong_id, tinh_trang_hn)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        mid,
        data.get("ho_ten", ""),
        data.get("gioi_tinh", "Nam"),
        int(data.get("the_he", 1)) if data.get("the_he") else None,
        data.get("vai_ve", ""),
        data.get("sinh", ""),
        data.get("mat_duong", ""),
        data.get("mat_am", ""),
        str(data.get("huong_tho", "")),
        data.get("que_quan", ""),
        data.get("phan_mo", ""),
        data.get("ghi_chu", ""),
        data.get("cha_id", ""),
        data.get("me_id", ""),
        data.get("vo_chong_id", ""),
        data.get("tinh_trang_hn", "")
    ))
    conn.commit()
    conn.close()
    sync_to_json_and_excel()
    return get_member_by_id(mid)

def update_member(member_id, data):
    conn = get_db()
    c = conn.cursor()
    c.execute("""
        UPDATE thanh_vien SET
            ho_ten = ?,
            gioi_tinh = ?,
            the_he = ?,
            vai_ve = ?,
            sinh = ?,
            mat_duong = ?,
            mat_am = ?,
            huong_tho = ?,
            que_quan = ?,
            phan_mo = ?,
            ghi_chu = ?,
            cha_id = ?,
            me_id = ?,
            vo_chong_id = ?,
            tinh_trang_hn = ?
        WHERE id = ?
    """, (
        data.get("ho_ten", ""),
        data.get("gioi_tinh", "Nam"),
        int(data.get("the_he", 1)) if data.get("the_he") else None,
        data.get("vai_ve", ""),
        data.get("sinh", ""),
        data.get("mat_duong", ""),
        data.get("mat_am", ""),
        str(data.get("huong_tho", "")),
        data.get("que_quan", ""),
        data.get("phan_mo", ""),
        data.get("ghi_chu", ""),
        data.get("cha_id", ""),
        data.get("me_id", ""),
        data.get("vo_chong_id", ""),
        data.get("tinh_trang_hn", ""),
        member_id
    ))
    conn.commit()
    conn.close()
    sync_to_json_and_excel()
    return get_member_by_id(member_id)

def delete_member(member_id):
    conn = get_db()
    c = conn.cursor()
    c.execute("DELETE FROM thanh_vien WHERE id = ?", (member_id,))
    conn.commit()
    conn.close()
    sync_to_json_and_excel()
    return True

def get_anniversaries():
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT * FROM ngay_gio ORDER BY id ASC")
    rows = [dict(r) for r in c.fetchall()]
    conn.close()
    return rows

def sync_to_json_and_excel():
    """Tự động đồng bộ CSDL SQLite ra file JSON và file Excel để dữ liệu luôn nhất quán"""
    try:
        members = get_all_members()
        anniversaries = get_anniversaries()

        # 1. Update JSON
        if os.path.exists(JSON_PATH):
            with open(JSON_PATH, 'r', encoding='utf-8') as f:
                jdata = json.load(f)
        else:
            jdata = {}
        jdata["thanh_vien"] = members
        jdata["ngay_gio"] = anniversaries
        with open(JSON_PATH, 'w', encoding='utf-8') as f:
            json.dump(jdata, f, ensure_ascii=False, indent=2)

        # 2. Update Excel Sheet 1
        if HAS_OPENPYXL and os.path.exists(EXCEL_PATH):
            wb = openpyxl.load_workbook(EXCEL_PATH)
            if "Danh sách Thành viên" in wb.sheetnames:
                ws = wb["Danh sách Thành viên"]
                # Clear existing rows from row 3
                max_r = ws.max_row
                if max_r > 2:
                    ws.delete_rows(3, max_r - 2)
                thin_border = Border(
                    left=Side(style='thin', color='D9D9D9'),
                    right=Side(style='thin', color='D9D9D9'),
                    top=Side(style='thin', color='D9D9D9'),
                    bottom=Side(style='thin', color='D9D9D9')
                )
                regular_font = Font(name="Arial", size=10)
                for m in members:
                    ws.append([
                        m["id"], m["ho_ten"], m["gioi_tinh"], m["the_he"], m["vai_ve"],
                        m["sinh"], m["mat_duong"], m["mat_am"], m["huong_tho"],
                        m["phan_mo"], m["ghi_chu"]
                    ])
                    r_idx = ws.max_row
                    ws.row_dimensions[r_idx].height = 22
                    for col_idx in range(1, 12):
                        cell = ws.cell(row=r_idx, column=col_idx)
                        cell.font = regular_font
                        cell.border = thin_border
                        if col_idx in [1, 3, 4, 6, 7, 8, 9]:
                            cell.alignment = Alignment(horizontal="center", vertical="center")
                        else:
                            cell.alignment = Alignment(vertical="center")
                wb.save(EXCEL_PATH)
    except Exception as e:
        print(f"[CẢNH BÁO ĐỒNG BỘ]: {e}")
